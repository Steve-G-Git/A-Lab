from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "tmp" / "resume" / "Steve_Garnet_Public_Resume.docx"


def set_cell_margins(cell, top=30, start=30, bottom=30, end=30):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, label, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "175A8B")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_separator(paragraph):
    run = paragraph.add_run("  |  ")
    run.font.color.rgb = RGBColor(110, 110, 110)


def add_section_heading(document, text):
    paragraph = document.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(1.5)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_role(document, title, organization, location, dates):
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(5.75)
    table.columns[1].width = Inches(1.45)
    table.rows[0].cells[0].width = Inches(5.75)
    table.rows[0].cells[1].width = Inches(1.45)
    table.allow_autofit = False
    for cell in table.rows[0].cells:
        set_cell_margins(cell, 0, 0, 0, 0)
        cell.vertical_alignment = 1

    left = table.rows[0].cells[0].paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    left.paragraph_format.keep_with_next = True
    title_run = left.add_run(title)
    title_run.bold = True
    role_details = [organization]
    if location:
        role_details.append(location)
    left.add_run(" | " + " | ".join(role_details))

    right = table.rows[0].cells[1].paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    right.paragraph_format.keep_with_next = True
    right_run = right.add_run(dates)
    right_run.bold = True
    return table


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.13)
    paragraph.paragraph_format.space_after = Pt(1.2)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.add_run(text)


def add_skill_line(document, label, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0.8)
    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    paragraph.add_run(text)


def build_resume():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.6)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(1.5)
    normal.paragraph_format.line_spacing = 1.0

    title_style = document.styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(23)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style_ppr = title_style.element.get_or_add_pPr()
    title_border = title_style_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_style_ppr.remove(title_border)

    heading_style = document.styles["Heading 1"]
    heading_style.font.name = "Arial"
    heading_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    heading_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    heading_style.font.color.rgb = RGBColor(0, 0, 0)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(0)
    title.add_run("Steve Garnet")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(2)
    subtitle_run = subtitle.add_run("IT SUPPORT  |  NETWORKING  |  TECHNICAL OPERATIONS")
    subtitle_run.bold = True
    subtitle_run.font.size = Pt(10)

    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(4)
    contact.add_run("Tooele, Utah")
    add_separator(contact)
    add_hyperlink(contact, "stevegarnet@outlook.com", "mailto:stevegarnet@outlook.com")
    add_separator(contact)
    add_hyperlink(contact, "LinkedIn", "https://www.linkedin.com/in/steve-garnet-502b10334/")
    add_separator(contact)
    add_hyperlink(contact, "GitHub", "https://github.com/Steve-G-Git")
    add_separator(contact)
    add_hyperlink(contact, "Portfolio", "https://steve-g-git.github.io/A-Lab/")

    add_section_heading(document, "Professional Summary")
    summary = document.add_paragraph()
    summary.paragraph_format.space_after = Pt(1.5)
    summary.add_run(
        "Entry-level IT support candidate with hands-on experience in Windows and Ubuntu Linux, computer hardware, "
        "networking, virtualization, troubleshooting, and technical documentation. Brings more than 20 years of "
        "manufacturing and operations experience, including team leadership, equipment troubleshooting, training, "
        "controlled procedures, and continuous improvement."
    )

    add_section_heading(document, "Technical Skills")
    add_skill_line(document, "Systems", "Windows 10/11, Ubuntu Server, Linux command line, VirtualBox, systemd, OpenSSH, UFW, Samba")
    add_skill_line(document, "Networking", "TCP/IP, IPv4, DNS, DHCP, NAT, subnetting, routing fundamentals, Wi-Fi, ping, ipconfig, tracert, nslookup")
    add_skill_line(document, "Hardware", "PC assembly, RAM and storage replacement, BIOS/UEFI, bootable media, peripherals, Cat5e/Cat6 termination and testing")
    add_skill_line(document, "Tools and Methods", "Git, GitHub, Bash, Python fundamentals, PowerShell fundamentals, runbooks, PDCA, root-cause analysis")

    add_section_heading(document, "Technical Projects")
    add_role(document, "LeanOps Lab", "Self-Directed", "Tooele, Utah", "2026")
    add_bullet(document, "Built an Ubuntu Server lab with NAT and isolated host-only networking, static IPv4 addressing, key-based SSH, and source-restricted firewall rules.")
    add_bullet(document, "Created Bash and Python health-monitoring workflows with systemd scheduling, durable event state, evidence collection, log retention, and grouped email notifications.")
    add_bullet(document, "Implemented role-based Samba shares, protected configuration backups, daily share-data backups, isolated restore tests, operational runbooks, and sanitized public evidence across 13 closed PDCA cycles.")

    add_role(document, "IT Support Portfolio", "Self-Directed", "Tooele, Utah", "2025 to Present")
    add_bullet(document, "Created and published networking study tools, an interactive OSI reference, and documented troubleshooting case studies using HTML, CSS, JavaScript, React, Git, and GitHub Pages.")
    add_bullet(document, "Recovered communication with a 3D-printer MCU by isolating the Linux host, USB serial interface, firmware, and controller layers; documented the failed attempts, recovery, and verification.")

    add_section_heading(document, "Professional Experience")
    add_role(document, "Team Lead", "US Synthetic", "Orem, Utah", "2017 to 2025")
    add_bullet(document, "Led 8 to 12 employees in precision manufacturing, coordinating priorities, training, quality, safety, documentation, and shift handoffs.")
    add_bullet(document, "Diagnosed equipment and process failures, documented findings, and coordinated corrective action with operators, maintenance, and leadership.")
    add_bullet(document, "Trained and certified employees using the EDGE method while reinforcing standardized work, 5S, PPE, and quality requirements.")

    add_role(document, "Sanitation Manager", "SupraNaturals", "", "2010 to 2017")
    add_bullet(document, "Managed industrial sanitation work, employee assignments, safety procedures, documentation, and production-area readiness.")

    add_role(document, "Line Lead", "Campbell Soup Company Ltd.", "", "2003 to 2010")
    add_bullet(document, "Led food-production work involving batching, production and packaging equipment, HMI controls, inventory systems, traceability, quality checks, and production reporting.")

    add_section_heading(document, "Education and Development")
    education = document.add_paragraph()
    education.paragraph_format.space_after = Pt(0)
    education.add_run("High School Diploma").bold = True
    education.add_run(" | CompTIA Network+ preparation through coursework and hands-on labs, in progress")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = "Steve Garnet Public Resume"
    document.core_properties.subject = "IT support and networking resume"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_resume()
